import pandas as pd, numpy as np, warnings; warnings.filterwarnings('ignore')
exec(open('metr.py').read().split("R=pd.read_pickle")[0])
U='/tmp/claude-0/-home-user-Minimo-Camino-Electrico---Balance-de-Potencia/aef6561c-d05e-5d6f-b069-1de59162b42e/scratchpad/up/bal.xlsm'
OUT='/home/user/Minimo-Camino-Electrico---Balance-de-Potencia/docs/precio_por_barra_criterios.xlsx'
f=pd.read_excel(U,sheet_name='05. BALANCE FÍSICO',header=None).iloc[6:][[7,13,14]]; f.columns=['barra','mw','tipo']
f['mw']=pd.to_numeric(f.mw,errors='coerce').fillna(0)
g=f[f.tipo!='Tx']; tx=f[f.tipo=='Tx']
iny=g[g.mw>0].groupby('barra').mw.sum(); ret=-g[g.mw<0].groupby('barra').mw.sum()
txmw=tx.groupby('barra').mw.apply(lambda s:s.abs().sum()/2); psuf=set(tx.barra)
p3=pd.read_excel(U,sheet_name='03. PRECIOS DE NUDO',header=None,usecols='B:D').iloc[4:]; p3.columns=['f','t','precio']
precio=p3[p3['f']=='PN2604'].set_index('t')['precio'].astype(float)
R=pd.read_pickle('banco.pkl'); R=R[R.actual.notna()]
nom={'P_ref220':'Referido','P_C1':'Referido + Subestación','actual':'Módulo actual'}
def tabla(k):
    d=R[R[k]!=R.PN2604].copy()
    o=pd.DataFrame({'Barra':d.barra,'Troncal PN2604':d.PN2604,f'Troncal {nom[k]}':d[k],
        'Precio PN2604 ($/kW/mes)':d.PN2604.map(precio).round(1),f'Precio {nom[k]} ($/kW/mes)':d[k].map(precio).round(1)})
    o['Δ precio ($/kW/mes)']=(o.iloc[:,4]-o.iloc[:,3]).round(1); o['Δ precio (%)']=(100*o['Δ precio ($/kW/mes)']/o.iloc[:,3]).round(1)
    o['km entre troncales']=[round(km(a,b)) if np.isfinite(km(a,b)) else None for a,b in zip(d.PN2604,d[k])]
    o['Tiene PSUF (filas Tx)']=np.where(d.barra.isin(psuf),'Sí','No')
    o['En hoja 05']=np.where(d.barra.isin(set(f.barra)),'Sí','No')
    o['Inyección (MW)']=d.barra.map(iny).fillna(0).round(2).values; o['Retiro (MW)']=d.barra.map(ret).fillna(0).round(2).values
    o['Flujo Tx (MW)']=d.barra.map(txmw).fillna(0).round(1).values
    o['Monto inyección+retiro (MM$/mes)']=(o['Δ precio ($/kW/mes)'].abs()*(o['Inyección (MW)']+o['Retiro (MW)'])/1000).round(2)
    return o.sort_values('Δ precio ($/kW/mes)',key=abs,ascending=False)
hojas={nom[k].replace('Referido + Subestación','Ref+Subestación')+' vs PN2604':tabla(k) for k in ['P_C1','P_ref220','actual']}
res=[]
for n,o in hojas.items():
    for ps in ['Sí','No']:
        s=o[o['Tiene PSUF (filas Tx)']==ps]
        res.append({'Comparación':n,'Tiene PSUF':ps,'Barras':len(s),'|Δ precio| mediana':s['Δ precio ($/kW/mes)'].abs().median().round(0),
                    '|Δ precio| > 500':int((s['Δ precio ($/kW/mes)'].abs()>500).sum()),'Con inyección/retiro':int(((s['Inyección (MW)']+s['Retiro (MW)'])>0).sum()),
                    'Monto (MM$/mes)':round(s['Monto inyección+retiro (MM$/mes)'].sum(),1)})
res=pd.DataFrame(res)
with pd.ExcelWriter(OUT,engine='openpyxl') as w:
    pd.DataFrame({'Nota':['Precios: PN2604 (hoja 03 Balance_2609_def). MW: hoja 05, Cons Estimado 2026.',
        'Monto = |Δ precio| × (inyección + retiro) / 1000, en MM$/mes. Solo filas Gx/Bfact; las filas Tx no se valorizan aquí.',
        'Tiene PSUF = la barra tiene filas Tx en la hoja 05 (misma definición que mb_dim_barra del modelo_balance).',
        'Δ precio = precio de la troncal del criterio − precio de la troncal PN2604.']}).to_excel(w,sheet_name='Resumen',index=False,startrow=0)
    res.to_excel(w,sheet_name='Resumen',index=False,startrow=6)
    for n,o in hojas.items(): o.to_excel(w,sheet_name=n[:31],index=False)
    o=hojas['Ref+Subestación vs PN2604'].copy(); o['Cambio']=o['Troncal PN2604'].str.replace('_','').str[:-3]+' → '+o.iloc[:,2].str.replace('_','').str[:-3]
    gg=o.groupby('Cambio').agg(**{'Barras':('Barra','size'),'Con PSUF':('Tiene PSUF (filas Tx)',lambda s:(s=='Sí').sum()),'Δ precio ($/kW/mes)':('Δ precio ($/kW/mes)','first'),
        'km entre troncales':('km entre troncales','first'),'Inyección (MW)':('Inyección (MW)','sum'),'Retiro (MW)':('Retiro (MW)','sum'),'Flujo Tx (MW)':('Flujo Tx (MW)','sum'),
        'Monto iny+ret (MM$/mes)':('Monto inyección+retiro (MM$/mes)','sum')})
    gg['Monto si Tx se valorizara (MM$/mes)']=(gg['Δ precio ($/kW/mes)'].abs()*gg['Flujo Tx (MW)']/1000).round(1)
    gg.sort_values('Monto iny+ret (MM$/mes)',ascending=False).round(1).to_excel(w,sheet_name='Ref+Sub por cambio')
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
wb=load_workbook(OUT)
for ws in wb.worksheets:
    for c in ws[1 if ws.title!='Resumen' else 7]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='30849B')
    for col in ws.columns: ws.column_dimensions[col[0].column_letter].width=max(12,min(45,max(len(str(c.value or '')) for c in col)+2))
    ws.freeze_panes='B2' if ws.title!='Resumen' else None
wb['Resumen'].column_dimensions['A'].width=110
wb.save(OUT)
print(res.to_string(index=False))
print(hojas['Ref+Subestación vs PN2604'].head(5).to_string(index=False))
